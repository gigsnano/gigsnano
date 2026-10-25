from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook

from gigsnano_api.features.work_item_io.columns import (
    COLUMNS,
    SHEET_NAME,
    Column,
    columns_by_key,
    header_map,
)


@dataclass
class ParsedRow:
    row: int
    values: dict[str, object] = field(default_factory=dict)   # field-columns, coerced
    raw: dict[str, str] = field(default_factory=dict)          # every present column key -> raw str
    errors: list[str] = field(default_factory=list)


def _cell_str(c: object) -> str:
    if c is None:
        return ""
    if isinstance(c, datetime):
        return c.date().isoformat()
    if isinstance(c, date):
        return c.isoformat()
    return str(c)


def _rows_from_bytes(content: bytes, fmt: str) -> list[list[str]]:
    if fmt == "csv":
        text = content.decode("utf-8-sig", errors="replace")
        return [list(r) for r in csv.reader(io.StringIO(text))]
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as err:  # noqa: BLE001
        raise ValueError("unreadable file") from err
    # Fall back to the active sheet when SHEET_NAME isn't present, so a
    # hand-made single-sheet file still parses.
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    if ws is None:
        raise ValueError("unreadable file")
    return [[_cell_str(c) for c in row]
            for row in ws.iter_rows(values_only=True)]


def parse_rows(
    content: bytes, fmt: str, columns: list[Column] | None = None,
) -> list[ParsedRow]:
    cols = columns or COLUMNS
    hmap = header_map(cols)
    by_key = columns_by_key(cols)
    grid = _rows_from_bytes(content, fmt)
    if not grid:
        return []
    header_keys = [hmap.get(str(h).strip().lower()) for h in grid[0]]
    out: list[ParsedRow] = []
    for i, raw_row in enumerate(grid[1:], start=2):
        cells: dict[str, str] = {
            k: (raw_row[j] if j < len(raw_row) else "")
            for j in range(len(header_keys))
            if (k := header_keys[j]) is not None
        }
        if not any((v or "").strip() for v in cells.values()):
            continue  # fully blank row
        pr = ParsedRow(row=i, raw={k: (v or "").strip() for k, v in cells.items()})
        if not pr.raw.get("title"):
            pr.errors.append("missing title")
        for key, raw in pr.raw.items():
            col = by_key[key]
            if col.coerce and raw != "":
                try:
                    pr.values[key] = col.coerce(raw)
                except ValueError as err:
                    # Wrap error with column context
                    err_msg = str(err)
                    # Include column key to ensure contextual keywords appear
                    pr.errors.append(f"{key}: {err_msg}")
        out.append(pr)
    return out
