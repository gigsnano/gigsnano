from io import BytesIO

from fastapi import HTTPException, status
from openpyxl import load_workbook

from gigsnano_api.features.work_item_io.columns import (
    COLUMNS,
    DEFAULT_KEYS,
    HEADER_TO_KEY,
    RenderCtx,
    coerce_date,
    coerce_estimate,
    coerce_priority,
    coerce_status,
    coerce_type,
)
from gigsnano_api.features.work_item_io.render import render_csv, render_xlsx
from gigsnano_api.models import EStoryPriority, EStoryStatus, EStoryType, Story


def test_registry_order_and_headers():
    assert [c.key for c in COLUMNS] == [
        "code", "title", "type", "parent", "status", "priority",
        "assignee", "labels", "start_date", "target_date", "estimate", "description",
    ]
    assert HEADER_TO_KEY["status"] == "status"
    assert HEADER_TO_KEY["start date"] == "start_date"
    assert DEFAULT_KEYS == [c.key for c in COLUMNS]


def test_coercers_accept_valid_and_reject_invalid():
    assert coerce_status("In Progress") == EStoryStatus.IN_PROGRESS
    assert coerce_type("epic") == EStoryType.EPIC
    assert coerce_priority("None") == EStoryPriority.NONE
    assert coerce_date("2026-07-14").isoformat() == "2026-07-14"
    assert coerce_estimate("5") == 5
    for bad in (coerce_status, coerce_type, coerce_priority):
        try:
            bad("nonsense")
            assert False
        except ValueError:
            pass
    for bad_val in ("2026-13-99", "not-a-date"):
        try:
            coerce_date(bad_val)
            assert False
        except ValueError:
            pass
    try:
        coerce_estimate("3.5")
        assert False
    except ValueError:
        pass


# A2: Workbook renderer tests


def test_render_xlsx_selected_columns_in_order():
    def _story(**kw):
        s = Story(project_id="p1", title=kw.get("title", "T"))
        for k, v in kw.items():
            setattr(s, k, v)
        return s

    ctx = RenderCtx(identifier="WEB", email_of={"u1": "a@x.com"},
                    code_of={"s1": "WEB-1"}, label_name_of={})
    s = _story()
    s.id = "s1"
    s.assignee_id = "u1"
    s.type = EStoryType.EPIC
    data = render_xlsx([s], ["code", "title", "type", "assignee"], ctx)
    wb = load_workbook(BytesIO(data))
    ws = wb["Work Items"]
    assert [c.value for c in ws[1]] == ["Code", "Title", "Type", "Assignee"]
    assert [c.value for c in ws[2]] == ["WEB-1", "T", "epic", "a@x.com"]


def test_render_xlsx_orders_by_registry_not_keys():
    def _story(**kw):
        s = Story(project_id="p1", title=kw.get("title", "T"))
        for k, v in kw.items():
            setattr(s, k, v)
        return s

    ctx = RenderCtx(identifier="WEB", email_of={},
                    code_of={"s1": "WEB-1"}, label_name_of={})
    s = _story()
    s.id = "s1"
    data = render_xlsx([s], ["assignee", "code"], ctx)
    wb = load_workbook(BytesIO(data))
    ws = wb["Work Items"]
    # Registry order: code, title, ..., assignee
    # Keys are: ["assignee", "code"] (reversed)
    # Expected output in registry order: ["Code", "Assignee"]
    assert [c.value for c in ws[1]] == ["Code", "Assignee"]
    # Empty values in Excel become None
    assert [c.value for c in ws[2]] == ["WEB-1", None]


def test_render_csv_selected_columns():
    def _story(**kw):
        s = Story(project_id="p1", title=kw.get("title", "T"))
        for k, v in kw.items():
            setattr(s, k, v)
        return s

    ctx = RenderCtx(identifier="WEB", email_of={},
                    code_of={"s1": "WEB-1"}, label_name_of={})
    s = _story(title="My Story")
    s.id = "s1"
    csv_data = render_csv([s], ["code", "title"], ctx)
    lines = [line.strip() for line in csv_data.strip().split("\n")]
    assert lines[0] == "Code,Title"
    assert lines[1] == "WEB-1,My Story"


# A3: Template builder tests


def test_template_has_headers_and_example_rows():
    from gigsnano_api.features.work_item_io.template import build_template

    wb = load_workbook(BytesIO(build_template()))
    ws = wb["Work Items"]
    headers = [c.value for c in ws[1]]
    assert headers[0] == "Code" and "Parent" in headers and "Assignee" in headers
    assert ws.max_row >= 3  # header + 2 example rows
    type_col = headers.index("Type")
    assert ws[2][type_col].value.lower() == "epic"
    assert ws[3][type_col].value.lower() == "story"


# A4: Parser tests


def _sheet_bytes(rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Work Items"
    headers = ["Code", "Title", "Type", "Parent", "Status", "Priority",
               "Assignee", "Labels", "Start date", "Target date", "Estimate", "Description"]
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_valid_and_error_rows():
    from gigsnano_api.features.work_item_io.parse import parse_rows

    data = _sheet_bytes([
        {"Title": "Epic A", "Type": "Epic", "Priority": "High"},
        {"Title": "", "Type": "Story"},                         # missing title
        {"Title": "Bad date", "Start date": "2026-99-99"},      # bad date
        {"Title": "Kid", "Parent": "Epic A", "Assignee": "a@x.com", "Labels": "ui, api"},
    ])
    rows = parse_rows(data, "xlsx")
    assert len(rows) == 4
    assert rows[0].values["type"].value == "epic"
    assert "missing title" in rows[1].errors[0]
    assert any("date" in e for e in rows[2].errors)
    assert rows[3].raw["parent"] == "Epic A"
    assert rows[3].raw["assignee"] == "a@x.com"
    assert rows[3].raw["labels"] == "ui, api"


def test_parse_reads_native_excel_date():
    """openpyxl returns real datetime.date/datetime.datetime objects for
    date-typed cells (not strings) -- the parser must not blanket-str() them."""
    from datetime import date

    from openpyxl import Workbook

    from gigsnano_api.features.work_item_io.parse import parse_rows

    wb = Workbook()
    ws = wb.active
    ws.title = "Work Items"
    headers = ["Code", "Title", "Type", "Parent", "Status", "Priority",
               "Assignee", "Labels", "Start date", "Target date", "Estimate", "Description"]
    ws.append(headers)
    ws.append(["", "Native Date Story", "", "", "", "", "", "", date(2026, 7, 14), "", "", ""])
    buf = BytesIO()
    wb.save(buf)

    rows = parse_rows(buf.getvalue(), "xlsx")
    assert len(rows) == 1
    assert rows[0].errors == []
    assert rows[0].values["start_date"].isoformat() == "2026-07-14"


def test_parse_csv_branch():
    from gigsnano_api.features.work_item_io.parse import parse_rows

    csv_bytes = (
        "Code,Title,Type,Parent,Status,Priority,Assignee,Labels,"
        "Start date,Target date,Estimate,Description\n"
        ",CSV Story,epic,,,,,,,,,\n"
    ).encode("utf-8")
    rows = parse_rows(csv_bytes, "csv")
    assert len(rows) == 1
    assert rows[0].raw["title"] == "CSV Story"
    assert rows[0].values["type"].value == "epic"


def test_parse_unreadable_file_raises():
    import pytest

    from gigsnano_api.features.work_item_io.parse import parse_rows

    with pytest.raises(ValueError):
        parse_rows(b"not a workbook", "xlsx")


# A5: Model project-scope columns


def test_job_models_have_project_scope_columns():
    from gigsnano_api.models import ExportJob, ImportJob

    e = ExportJob(workspace_id="w1", name="n", format="xlsx",
                  project_id="p1", columns=["code", "title"], filters={"status_in": ["todo"]})
    assert e.project_id == "p1" and e.columns == ["code", "title"]
    i = ImportJob(workspace_id="w1", source_format="xlsx", raw_content="", project_id="p1")
    assert i.project_id == "p1"


# A6: Two-pass idempotent importer + project export runner


async def _setup_workspace_member_project(db_session):
    """A workspace + an active member + a project in that same workspace --
    the importer resolves assignee emails against workspace members, so the
    member must share the project's workspace_id (make_project's own workspace
    doesn't have any members by default)."""
    from uuid import uuid4

    from gigsnano_api.auth.models import User
    from gigsnano_api.models import (
        EMemberStatus,
        EWorkspaceRole,
        Project,
        Workspace,
        WorkspaceMember,
    )

    ws = Workspace(slug=f"ws-{uuid4().hex[:8]}", name="WS")
    user = User(email="member@test.com", display_name="Member", password_hash="x")
    db_session.add_all([ws, user])
    await db_session.flush()
    db_session.add(WorkspaceMember(
        workspace_id=ws.id, user_id=user.id,
        role=EWorkspaceRole.MEMBER, status=EMemberStatus.ACTIVE,
    ))
    project = Project(
        workspace_id=ws.id, name="P", identifier="TST",
        cover_image_url="https://img.test/cover.jpg",
    )
    db_session.add(project)
    await db_session.commit()
    await db_session.refresh(project)
    await db_session.refresh(user)
    return ws, user, project


async def test_import_creates_epic_and_links_child(db_session):
    import base64

    from sqlalchemy import select

    from gigsnano_api.features.work_item_io.importer import run_work_item_import
    from gigsnano_api.models import EImportStatus, ImportJob, Story

    _ws, member_user, project = await _setup_workspace_member_project(db_session)
    content = _sheet_bytes([
        {"Title": "Epic A", "Type": "Epic"},
        {"Title": "Child 1", "Type": "Story", "Parent": "Epic A",
         "Assignee": member_user.email},
    ])
    job = ImportJob(
        workspace_id=project.workspace_id, source_format="xlsx",
        raw_content=base64.b64encode(content).decode(),
        project_id=project.id, summary={},
    )
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(job)

    await run_work_item_import(db_session, job.id)
    await db_session.refresh(job)

    assert job.status == EImportStatus.COMPLETED
    assert job.summary["created"] == 2
    stories = (await db_session.execute(
        select(Story).where(Story.project_id == project.id)
    )).scalars().all()
    epic = next(s for s in stories if s.title == "Epic A")
    child = next(s for s in stories if s.title == "Child 1")
    assert child.parent_id == epic.id
    assert child.assignee_id == member_user.id


async def test_round_trip_is_idempotent(db_session):
    import base64

    from sqlalchemy import func, select

    from gigsnano_api.features.work_item_io.columns import DEFAULT_KEYS
    from gigsnano_api.features.work_item_io.importer import (
        build_render_ctx,
        project_export_stories,
        run_work_item_import,
    )
    from gigsnano_api.features.work_item_io.render import render_xlsx
    from gigsnano_api.models import ImportJob, Story

    _ws, member_user, project = await _setup_workspace_member_project(db_session)
    s1 = Story(project_id=project.id, title="Story One", story_number=1, sort_order=0)
    s2 = Story(
        project_id=project.id, title="Story Two", story_number=2, sort_order=1,
        assignee_id=member_user.id,
    )
    db_session.add_all([s1, s2])
    await db_session.commit()

    ctx = await build_render_ctx(db_session, project)
    stories = await project_export_stories(db_session, project.id, None)
    data = render_xlsx(stories, DEFAULT_KEYS, ctx)

    job = ImportJob(
        workspace_id=project.workspace_id, source_format="xlsx",
        raw_content=base64.b64encode(data).decode(),
        project_id=project.id, summary={},
    )
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(job)

    before = len(stories)
    await run_work_item_import(db_session, job.id)
    await db_session.refresh(job)

    after = (await db_session.execute(
        select(func.count(Story.id)).where(Story.project_id == project.id)
    )).scalar_one()
    assert job.summary["created"] == 0
    assert job.summary["updated"] == before
    assert after == before  # no duplicates


async def test_unmatched_assignee_is_reported_not_fatal(db_session, make_project):
    import base64

    from sqlalchemy import select

    from gigsnano_api.features.work_item_io.importer import run_work_item_import
    from gigsnano_api.models import ImportJob, Story

    project = await make_project()
    content = _sheet_bytes([{"Title": "Orphan", "Assignee": "ghost@nowhere.com"}])
    job = ImportJob(
        workspace_id=project.workspace_id, source_format="xlsx",
        raw_content=base64.b64encode(content).decode(),
        project_id=project.id, summary={},
    )
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(job)

    await run_work_item_import(db_session, job.id)
    await db_session.refresh(job)

    assert "ghost@nowhere.com" in job.summary["unmatched_assignees"]
    assert job.summary["created"] == 1

    story = (await db_session.execute(
        select(Story).where(Story.project_id == project.id, Story.title == "Orphan")
    )).scalar_one()
    assert story.assignee_id is None


async def test_validate_is_dry_run(db_session, make_project):
    from sqlalchemy import func, select

    from gigsnano_api.features.work_item_io.importer import validate_import
    from gigsnano_api.models import Story

    project = await make_project()
    content = _sheet_bytes([
        {"Title": "New", "Type": "Story"},
        {"Title": "", "Type": "Story"},
    ])
    preview = await validate_import(db_session, project, content, "xlsx")
    assert preview["to_create"] == 1
    assert preview["errors"][0]["row"] == 3
    count = (await db_session.execute(
        select(func.count(Story.id)).where(Story.project_id == project.id)
    )).scalar_one()
    assert count == 0  # nothing written


async def test_import_crash_sets_failed(db_session, make_project, monkeypatch):
    """A raise anywhere in the pass-1/pass-2 upsert loop (e.g. `_insert_story`
    exhausting its story-number retries with an HTTPException(409)) must still
    leave the job in a terminal FAILED status with an error recorded --
    not stuck at IMPORTING forever."""
    import base64

    from gigsnano_api.features.work_item_io import importer as importer_mod
    from gigsnano_api.models import EImportStatus, ImportJob

    project = await make_project()
    content = _sheet_bytes([{"Title": "Will crash", "Type": "Story"}])
    job = ImportJob(
        workspace_id=project.workspace_id, source_format="xlsx",
        raw_content=base64.b64encode(content).decode(),
        project_id=project.id, summary={},
    )
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(job)

    async def _boom(*_args, **_kwargs):
        raise HTTPException(status.HTTP_409_CONFLICT, "Could not allocate a story number")

    monkeypatch.setattr(importer_mod, "_insert_story", _boom)

    await importer_mod.run_work_item_import(db_session, job.id)
    await db_session.refresh(job)

    assert job.status == EImportStatus.FAILED
    assert job.summary["errors"]
    assert job.summary["errors"][0]["reason"]


async def test_import_prepare_failure_sets_failed(db_session, make_project, monkeypatch):
    """A DB-layer failure in the PARSING->IMPORTING window (the `_prepare`
    SELECTs or the IMPORTING commit, both of which now live inside the main
    try) must still land the job at FAILED -- not stuck at PARSING/IMPORTING
    forever. Regression test for the finding: this window used to sit
    *before* the guarding try/except."""
    import base64

    from gigsnano_api.features.work_item_io import importer as importer_mod
    from gigsnano_api.models import EImportStatus, ImportJob

    project = await make_project()
    content = _sheet_bytes([{"Title": "Whatever", "Type": "Story"}])
    job = ImportJob(
        workspace_id=project.workspace_id, source_format="xlsx",
        raw_content=base64.b64encode(content).decode(),
        project_id=project.id, summary={},
    )
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(job)

    async def _boom(*_args, **_kwargs):
        raise RuntimeError("db exploded during prepare")

    monkeypatch.setattr(importer_mod, "_prepare", _boom)

    await importer_mod.run_work_item_import(db_session, job.id)
    await db_session.refresh(job)

    assert job.status == EImportStatus.FAILED
    assert job.summary["errors"]
    assert "db exploded during prepare" in job.summary["errors"][0]["reason"]


async def test_export_runner_renders_project(db_session, make_project):
    import base64
    from io import BytesIO

    from openpyxl import load_workbook

    from gigsnano_api.features.work_item_io.columns import DEFAULT_KEYS
    from gigsnano_api.features.work_item_io.importer import run_work_item_export
    from gigsnano_api.models import EExportStatus, ExportJob, Story

    project = await make_project()
    s1 = Story(project_id=project.id, title="Story One", story_number=1, sort_order=0)
    s2 = Story(project_id=project.id, title="Story Two", story_number=2, sort_order=1)
    bmad_card = Story(
        project_id=project.id, title="BMAD Card", story_number=3, sort_order=2,
        bmad_workflow="analyst-brief",
    )
    db_session.add_all([s1, s2, bmad_card])
    await db_session.commit()

    job = ExportJob(
        workspace_id=project.workspace_id, name="Export", format="xlsx",
        project_id=project.id, columns=DEFAULT_KEYS,
    )
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(job)

    await run_work_item_export(db_session, job.id)
    await db_session.refresh(job)

    assert job.status == EExportStatus.READY
    assert job.content

    wb = load_workbook(BytesIO(base64.b64decode(job.content)))
    ws = wb["Work Items"]
    title_col = [c.value for c in ws[1]].index("Title")
    titles = [row[title_col].value for row in ws.iter_rows(min_row=2)]
    assert "Story One" in titles
    assert "Story Two" in titles
    assert "BMAD Card" not in titles  # bmad_workflow rows are excluded from exports


async def test_export_setup_failure_sets_failed(db_session, make_project, monkeypatch):
    """A failure in export setup (building the render context, before the
    render itself) must still land the job at FAILED -- the whole body is
    now guarded, not just the render call. Regression test for the finding:
    `session.get`/setup used to sit outside the guarding try/except."""
    from gigsnano_api.features.work_item_io import importer as importer_mod
    from gigsnano_api.models import EExportStatus, ExportJob

    project = await make_project()
    job = ExportJob(
        workspace_id=project.workspace_id, name="Export", format="xlsx",
        project_id=project.id, columns=["code", "title"],
    )
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(job)

    async def _boom(*_args, **_kwargs):
        raise RuntimeError("db exploded building render ctx")

    monkeypatch.setattr(importer_mod, "build_render_ctx", _boom)

    await importer_mod.run_work_item_export(db_session, job.id)
    await db_session.refresh(job)

    assert job.status == EExportStatus.FAILED


# A7: HTTP routes + registration


async def _owner_project(client):
    await client.post("/api/auth/sign-up/", json={
        "display_name": "Ada", "email": "ada@example.test", "password": "s3cret!"})
    await client.post("/api/workspaces/", json={"name": "Acme", "slug": "acme"})
    project = (await client.post("/api/workspaces/acme/projects/", json={
        "name": "Mobile App", "identifier": "MBA", "flow": "fresh",
        "cover_image_url": "https://img.test/c.jpg"})).json()
    return project  # slug is "acme"


async def test_template_route_returns_xlsx(client):
    project = await _owner_project(client)
    r = await client.get(
        f"/api/workspaces/acme/projects/{project['id']}/work-items/export/template/")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]


async def test_create_export_returns_pending(client):
    project = await _owner_project(client)
    r = await client.post(
        f"/api/workspaces/acme/projects/{project['id']}/work-items/export/",
        json={"columns": ["code", "title"], "format": "xlsx", "scope": "all"})
    assert r.status_code == 201
    assert r.json()["status"] == "pending"


async def test_validate_route_dry_run(client):
    project = await _owner_project(client)
    data = _sheet_bytes([{"Title": "X", "Type": "Story"}])
    r = await client.post(
        f"/api/workspaces/acme/projects/{project['id']}/work-items/import/validate/",
        files={"file": ("wi.xlsx", data,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"format": "xlsx"})
    assert r.status_code == 200
    assert r.json()["to_create"] == 1


async def test_viewer_cannot_import(client):
    project = await _owner_project(client)
    grace = (await client.post("/api/auth/sign-up/", json={
        "display_name": "Grace Hopper", "email": "grace@example.test",
        "password": "s3cret!"})).json()
    await client.post("/api/auth/sign-out/")
    await client.post("/api/auth/sign-in/", json={
        "email": "ada@example.test", "password": "s3cret!"})
    await client.post("/api/workspaces/acme/members/", json={"user_id": grace["id"]})
    await client.post(
        f"/api/workspaces/acme/projects/{project['id']}/members/",
        json={"user_id": grace["id"], "role": "viewer"})
    await client.post("/api/auth/sign-out/")
    await client.post("/api/auth/sign-in/", json={
        "email": "grace@example.test", "password": "s3cret!"})

    data = _sheet_bytes([{"Title": "X"}])
    r = await client.post(
        f"/api/workspaces/acme/projects/{project['id']}/work-items/import/",
        files={"file": ("wi.xlsx", data, "application/octet-stream")},
        data={"format": "xlsx"})
    assert r.status_code == 403


# A7b: project-scoped import status route (member-pollable)


async def test_import_status_route_returns_job(client):
    project = await _owner_project(client)
    data = _sheet_bytes([{"Title": "X"}])
    created = (await client.post(
        f"/api/workspaces/acme/projects/{project['id']}/work-items/import/",
        files={"file": ("wi.xlsx", data,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"format": "xlsx"})).json()

    r = await client.get(
        f"/api/workspaces/acme/projects/{project['id']}/work-items/import/{created['id']}/")
    assert r.status_code == 200
    body = r.json()
    assert body["id"] == created["id"]
    assert "status" in body
    assert "summary" in body


async def test_import_status_404_for_other_project(client):
    project_a = await _owner_project(client)
    project_b = (await client.post("/api/workspaces/acme/projects/", json={
        "name": "Other App", "identifier": "OTH", "flow": "fresh",
        "cover_image_url": "https://img.test/c2.jpg"})).json()

    data = _sheet_bytes([{"Title": "X"}])
    created = (await client.post(
        f"/api/workspaces/acme/projects/{project_a['id']}/work-items/import/",
        files={"file": ("wi.xlsx", data,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"format": "xlsx"})).json()

    r = await client.get(
        f"/api/workspaces/acme/projects/{project_b['id']}/work-items/import/{created['id']}/")
    assert r.status_code == 404


async def test_export_status_route_returns_job(client):
    project = await _owner_project(client)
    created = (await client.post(
        f"/api/workspaces/acme/projects/{project['id']}/work-items/export/",
        json={"columns": ["code", "title"], "format": "xlsx", "scope": "all"})).json()

    r = await client.get(
        f"/api/workspaces/acme/projects/{project['id']}/work-items/export/{created['id']}/")
    assert r.status_code == 200
    body = r.json()
    assert body["id"] == created["id"]
    assert "status" in body
    assert "format" in body
    # BackgroundTask doesn't execute in the test client, so status may still
    # be "pending" -- we're only asserting the route returns the job.


async def test_export_status_404_for_other_project(client):
    project_a = await _owner_project(client)
    project_b = (await client.post("/api/workspaces/acme/projects/", json={
        "name": "Other App", "identifier": "OTH", "flow": "fresh",
        "cover_image_url": "https://img.test/c2.jpg"})).json()

    created = (await client.post(
        f"/api/workspaces/acme/projects/{project_a['id']}/work-items/export/",
        json={"columns": ["code", "title"], "format": "xlsx", "scope": "all"})).json()

    r = await client.get(
        f"/api/workspaces/acme/projects/{project_b['id']}/work-items/export/{created['id']}/")
    assert r.status_code == 404


# A1 (GIN-76 2b): project-aware column registry -- build_columns


from gigsnano_api.features.work_item_io.columns import build_columns, header_map  # noqa: E402
from gigsnano_api.models import ECustomFieldType, ProjectCustomField  # noqa: E402


def _cf(id, name, ftype, options=None):
    f = ProjectCustomField(project_id="p", name=name, field_type=ftype, options=options)
    f.id = id
    return f


def test_build_columns_adds_custom_and_validates():
    fields = [_cf("f1", "Team", ECustomFieldType.TEXT),
              _cf("f2", "Prio", ECustomFieldType.SELECT, ["Low", "High"])]
    cols = build_columns(fields)
    keys = [c.key for c in cols]
    assert "cf:f1" in keys and "cf:f2" in keys
    hm = header_map(cols)
    assert hm["team"] == "cf:f1" and hm["prio"] == "cf:f2"
    prio = next(c for c in cols if c.key == "cf:f2")
    assert prio.custom_field_id == "f2"
    assert prio.coerce("High") == "High"
    try:
        prio.coerce("Nope")
        assert False
    except ValueError:
        pass
    num = build_columns([_cf("f3", "Pts", ECustomFieldType.NUMBER)])[-1]
    assert num.coerce("8") == "8"
    try:
        num.coerce("abc")
        assert False
    except ValueError:
        pass


# A4 (GIN-76 2b): export route accepts cf:* custom-field column keys


async def test_export_route_accepts_custom_column(client, db_session):
    """The route used to filter payload.columns against the static
    DEFAULT_KEYS, silently dropping `cf:{id}` keys. It must now pass
    `cf:*` keys through -- the runner's own column resolution renders
    them (or drops truly invalid keys), so the route doesn't need to
    know about custom fields at all."""
    from gigsnano_api.features.work_item_io.importer import run_work_item_export

    project = await _owner_project(client)
    project_id = project["id"]

    field = (await client.post(
        f"/api/workspaces/acme/projects/{project_id}/custom-fields/",
        json={"name": "Team", "field_type": "text"})).json()
    field_id = field["id"]

    story = (await client.post(
        f"/api/workspaces/acme/projects/{project_id}/stories/",
        json={"title": "Story One"})).json()
    r = await client.patch(
        f"/api/workspaces/acme/projects/{project_id}/stories/{story['id']}/",
        json={"custom_fields": {field_id: "Platform"}})
    assert r.status_code == 200

    created = (await client.post(
        f"/api/workspaces/acme/projects/{project_id}/work-items/export/",
        json={"columns": ["code", "title", f"cf:{field_id}"],
              "format": "xlsx", "scope": "all"})).json()
    assert created["status"] == "pending"
    job_id = created["id"]

    status_url = f"/api/workspaces/acme/projects/{project_id}/work-items/export/{job_id}/"
    body = None
    for _ in range(10):
        body = (await client.get(status_url)).json()
        if body["status"] == "ready":
            break
    else:
        # BackgroundTask didn't execute inside the test ASGI transport (as
        # already noted by test_export_status_route_returns_job) -- fall
        # back to running the exporter directly against the shared
        # db_session (same underlying sqlite engine as `client`/`app`).
        await run_work_item_export(db_session, job_id)
        body = (await client.get(status_url)).json()

    assert body["status"] == "ready"
    assert body["download_url"]

    dl = await client.get(f"/api{body['download_url']}")
    assert dl.status_code == 200

    wb = load_workbook(BytesIO(dl.content))
    ws = wb["Work Items"]
    headers = [c.value for c in ws[1]]
    assert "Team" in headers
    team_col = headers.index("Team")
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert any(row[team_col] == "Platform" for row in rows)


def test_custom_export_reads_ctx():
    cols = build_columns([_cf("f1", "Team", ECustomFieldType.TEXT)])
    col = next(c for c in cols if c.key == "cf:f1")
    ctx = RenderCtx(identifier="P", email_of={}, code_of={}, label_name_of={},
                    custom_of={"s1": {"f1": "Platform"}})
    s = Story(project_id="p", title="t")
    s.id = "s1"
    assert col.export(s, ctx) == "Platform"


# A2 (GIN-76 2b): render/template/parse operate on a resolved column set


def test_render_xlsx_with_custom_columns():
    cols = build_columns([_cf("f1", "Team", ECustomFieldType.TEXT)])
    ctx = RenderCtx(identifier="WEB", email_of={}, code_of={"s1": "WEB-1"}, label_name_of={},
                     custom_of={"s1": {"f1": "Platform"}})
    s = Story(project_id="p1", title="My Story")
    s.id = "s1"
    data = render_xlsx([s], ["title", "cf:f1"], ctx, columns=cols)
    wb = load_workbook(BytesIO(data))
    ws = wb["Work Items"]
    assert [c.value for c in ws[1]] == ["Title", "Team"]
    assert [c.value for c in ws[2]] == ["My Story", "Platform"]


def test_render_csv_with_custom_columns():
    cols = build_columns([_cf("f1", "Team", ECustomFieldType.TEXT)])
    ctx = RenderCtx(identifier="WEB", email_of={}, code_of={"s1": "WEB-1"}, label_name_of={},
                     custom_of={"s1": {"f1": "Platform"}})
    s = Story(project_id="p1", title="My Story")
    s.id = "s1"
    csv_data = render_csv([s], ["title", "cf:f1"], ctx, columns=cols)
    lines = [line.strip() for line in csv_data.strip().split("\n")]
    assert lines[0] == "Title,Team"
    assert lines[1] == "My Story,Platform"


def test_build_template_with_custom_columns():
    from gigsnano_api.features.work_item_io.template import build_template

    cols = build_columns([_cf("f1", "Team", ECustomFieldType.TEXT)])
    wb = load_workbook(BytesIO(build_template(columns=cols)))
    ws = wb["Work Items"]
    headers = [c.value for c in ws[1]]
    assert "Team" in headers
    team_col = headers.index("Team")
    # Custom columns have no curated example value -> blank example cell.
    assert ws[2][team_col].value in (None, "")
    assert ws[3][team_col].value in (None, "")


def test_parse_rows_with_custom_columns():
    from openpyxl import Workbook

    from gigsnano_api.features.work_item_io.parse import parse_rows

    cols = build_columns([_cf("f1", "Team", ECustomFieldType.SELECT, ["Alpha", "Beta"])])
    wb = Workbook()
    ws = wb.active
    ws.title = "Work Items"
    ws.append(["Title", "Team"])
    ws.append(["Good Row", "Alpha"])
    ws.append(["Bad Row", "NotAnOption"])
    buf = BytesIO()
    wb.save(buf)

    rows = parse_rows(buf.getvalue(), "xlsx", columns=cols)
    assert len(rows) == 2
    assert rows[0].raw["cf:f1"] == "Alpha"
    assert rows[0].values["cf:f1"] == "Alpha"
    assert not rows[0].errors
    assert rows[1].raw["cf:f1"] == "NotAnOption"
    assert "cf:f1" not in rows[1].values
    assert any("cf:f1" in e for e in rows[1].errors)


# A3 (GIN-76 2b): importer resolves columns + applies custom values


async def test_export_custom_field_value(db_session, make_project):
    """Seeding a story + a custom field + a StoryCustomFieldValue and running
    the export must produce a workbook with the custom column header + value."""
    import base64

    from gigsnano_api.features.work_item_io.importer import run_work_item_export
    from gigsnano_api.models import (
        EExportStatus,
        ExportJob,
        Story,
        StoryCustomFieldValue,
    )

    project = await make_project()
    field = ProjectCustomField(project_id=project.id, name="Team", field_type=ECustomFieldType.TEXT)
    db_session.add(field)
    await db_session.flush()
    story = Story(project_id=project.id, title="Story One", story_number=1, sort_order=0)
    db_session.add(story)
    await db_session.flush()
    db_session.add(StoryCustomFieldValue(story_id=story.id, field_id=field.id, value="Platform"))
    await db_session.commit()

    job = ExportJob(
        workspace_id=project.workspace_id, name="Export", format="xlsx",
        project_id=project.id, columns=None,
    )
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(job)

    await run_work_item_export(db_session, job.id)
    await db_session.refresh(job)

    assert job.status == EExportStatus.READY
    wb = load_workbook(BytesIO(base64.b64decode(job.content)))
    ws = wb["Work Items"]
    headers = [c.value for c in ws[1]]
    assert "Team" in headers
    team_col = headers.index("Team")
    row_values = [row[team_col].value for row in ws.iter_rows(min_row=2)]
    assert "Platform" in row_values


async def test_round_trip_preserves_custom_value_no_dup(db_session, make_project):
    """Exporting a project with a custom value, then re-importing those exact
    bytes, must be a no-op on story count AND must not duplicate the
    StoryCustomFieldValue row for (story, field)."""
    import base64

    from sqlalchemy import select

    from gigsnano_api.features.work_item_io.importer import (
        build_render_ctx,
        project_export_stories,
        resolve_columns,
        run_work_item_import,
    )
    from gigsnano_api.features.work_item_io.render import render_xlsx
    from gigsnano_api.models import ImportJob, Story, StoryCustomFieldValue

    project = await make_project()
    field = ProjectCustomField(project_id=project.id, name="Team", field_type=ECustomFieldType.TEXT)
    db_session.add(field)
    await db_session.flush()
    story = Story(project_id=project.id, title="Story One", story_number=1, sort_order=0)
    db_session.add(story)
    await db_session.flush()
    db_session.add(StoryCustomFieldValue(story_id=story.id, field_id=field.id, value="Platform"))
    await db_session.commit()

    ctx = await build_render_ctx(db_session, project)
    cols = await resolve_columns(db_session, project)
    stories = await project_export_stories(db_session, project.id, None)
    data = render_xlsx(stories, [c.key for c in cols], ctx, columns=cols)

    job = ImportJob(
        workspace_id=project.workspace_id, source_format="xlsx",
        raw_content=base64.b64encode(data).decode(),
        project_id=project.id, summary={},
    )
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(job)

    await run_work_item_import(db_session, job.id)
    await db_session.refresh(job)

    assert job.summary["created"] == 0

    cfvs = (await db_session.execute(
        select(StoryCustomFieldValue).where(
            StoryCustomFieldValue.story_id == story.id, StoryCustomFieldValue.field_id == field.id,
        )
    )).scalars().all()
    assert len(cfvs) == 1
    assert cfvs[0].value == "Platform"


async def test_import_new_row_creates_custom_value(db_session, make_project):
    """A sheet with Title + a "Team" custom column must give the newly created
    story a StoryCustomFieldValue with the cell's value."""
    import base64

    from openpyxl import Workbook
    from sqlalchemy import select

    from gigsnano_api.features.work_item_io.importer import run_work_item_import
    from gigsnano_api.models import ImportJob, Story, StoryCustomFieldValue

    project = await make_project()
    field = ProjectCustomField(project_id=project.id, name="Team", field_type=ECustomFieldType.TEXT)
    db_session.add(field)
    await db_session.commit()
    await db_session.refresh(field)

    wb = Workbook()
    ws = wb.active
    ws.title = "Work Items"
    ws.append(["Title", "Team"])
    ws.append(["New Story", "Platform"])
    buf = BytesIO()
    wb.save(buf)

    job = ImportJob(
        workspace_id=project.workspace_id, source_format="xlsx",
        raw_content=base64.b64encode(buf.getvalue()).decode(),
        project_id=project.id, summary={},
    )
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(job)

    await run_work_item_import(db_session, job.id)
    await db_session.refresh(job)

    assert job.summary["created"] == 1
    story = (await db_session.execute(
        select(Story).where(Story.project_id == project.id, Story.title == "New Story")
    )).scalar_one()
    cfv = (await db_session.execute(
        select(StoryCustomFieldValue).where(
            StoryCustomFieldValue.story_id == story.id, StoryCustomFieldValue.field_id == field.id,
        )
    )).scalar_one()
    assert cfv.value == "Platform"


async def test_import_bad_custom_value_skips_row(db_session, make_project):
    """A SELECT custom column with a non-option value must error/skip that
    row while a valid row in the same file still imports."""
    import base64

    from openpyxl import Workbook
    from sqlalchemy import select

    from gigsnano_api.features.work_item_io.importer import run_work_item_import
    from gigsnano_api.models import ImportJob, Story, StoryCustomFieldValue

    project = await make_project()
    field = ProjectCustomField(
        project_id=project.id, name="Team", field_type=ECustomFieldType.SELECT,
        options=["Alpha", "Beta"],
    )
    db_session.add(field)
    await db_session.commit()
    await db_session.refresh(field)

    wb = Workbook()
    ws = wb.active
    ws.title = "Work Items"
    ws.append(["Title", "Team"])
    ws.append(["Good Row", "Alpha"])
    ws.append(["Bad Row", "NotAnOption"])
    buf = BytesIO()
    wb.save(buf)

    job = ImportJob(
        workspace_id=project.workspace_id, source_format="xlsx",
        raw_content=base64.b64encode(buf.getvalue()).decode(),
        project_id=project.id, summary={},
    )
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(job)

    await run_work_item_import(db_session, job.id)
    await db_session.refresh(job)

    assert job.summary["created"] == 1
    assert job.summary["skipped"] == 1
    assert any("cf:" in e["reason"] for e in job.summary["errors"])

    good = (await db_session.execute(
        select(Story).where(Story.project_id == project.id, Story.title == "Good Row")
    )).scalar_one()
    cfv = (await db_session.execute(
        select(StoryCustomFieldValue).where(
            StoryCustomFieldValue.story_id == good.id, StoryCustomFieldValue.field_id == field.id,
        )
    )).scalar_one()
    assert cfv.value == "Alpha"

    bad_rows = (await db_session.execute(
        select(Story).where(Story.project_id == project.id, Story.title == "Bad Row")
    )).scalars().all()
    assert bad_rows == []
