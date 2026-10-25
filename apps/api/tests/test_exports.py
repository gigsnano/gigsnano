import base64
from datetime import timedelta, timezone
from io import BytesIO

from openpyxl import load_workbook

from gigsnano_api.core.db import utcnow
from gigsnano_api.email.provider import FakeEmailProvider
from gigsnano_api.features.exports.router import run_export_in_fresh_session
from gigsnano_api.models import EExportStatus, ExportJob

CREATE_PROJECT = {
    "name": "Mobile App", "identifier": "MBA", "flow": "fresh", "has_ui": False,
    "cover_image_url": "https://img.test/c.jpg",
}


async def setup(client):
    await client.post("/api/auth/sign-up/", json={
        "display_name": "Ada", "email": "ada@example.test", "password": "s3cret!",
    })
    await client.post("/api/workspaces/", json={"name": "Acme", "slug": "acme"})
    p = (await client.post("/api/workspaces/acme/projects/", json=CREATE_PROJECT)).json()
    await client.post(f"/api/workspaces/acme/projects/{p['id']}/stories/", json={"title": "Task A"})
    return p


async def test_export_job_model_roundtrips_expires_at(db_session, make_workspace):
    ws = await make_workspace()
    expires = utcnow() + timedelta(days=7)
    job = ExportJob(workspace_id=ws.id, name="Backlog", format="csv", expires_at=expires)
    db_session.add(job)
    await db_session.commit()
    job_id = job.id

    db_session.expire_all()
    reloaded = await db_session.get(ExportJob, job_id)
    assert reloaded.expires_at is not None
    reloaded_expires = reloaded.expires_at
    if reloaded_expires.tzinfo is None:  # sqlite drops tzinfo on round-trip
        reloaded_expires = reloaded_expires.replace(tzinfo=timezone.utc)
    assert abs((reloaded_expires - expires).total_seconds()) < 1


async def _run_job(db_session, job_id: str, project_ids=None) -> None:
    """Drive the background render directly (deterministic in tests — see the
    gigsnano-api-bg-session-prod-db memory on why `run_export_in_fresh_session`
    binds to the request session's engine rather than the module async_session())."""
    await run_export_in_fresh_session(db_session.bind, job_id, project_ids, FakeEmailProvider())


async def test_create_export_is_pending_then_ready_after_running_job(app, client, db_session):
    await setup(client)
    res = await client.post("/api/workspaces/acme/exports/", json={"name": "Backlog", "format": "csv"})
    assert res.status_code == 201
    job = res.json()
    assert job["status"] == "pending"
    assert job["download_url"] is None
    assert job["expires_at"] is not None

    await _run_job(db_session, job["id"])

    listed = (await client.get("/api/workspaces/acme/exports/")).json()["items"]
    ready = next(j for j in listed if j["id"] == job["id"])
    assert ready["status"] == "ready"
    assert ready["download_url"] == f"/workspaces/acme/exports/{job['id']}/download/"

    dl = await client.get(f"/api/workspaces/acme/exports/{job['id']}/download/")
    assert dl.status_code == 200
    body = dl.text
    assert "id,title,status,priority" in body  # header row
    assert "Task A" in body  # one row per story


async def test_create_json_export(app, client, db_session):
    await setup(client)
    job = (await client.post(
        "/api/workspaces/acme/exports/", json={"name": "Data", "format": "json"})).json()
    await _run_job(db_session, job["id"])

    dl = await client.get(f"/api/workspaces/acme/exports/{job['id']}/download/")
    assert dl.status_code == 200
    import json as _json
    rows = _json.loads(dl.text)
    assert any(r["title"] == "Task A" for r in rows)


async def test_create_xlsx_export_content_is_reopenable_base64(app, client, db_session):
    await setup(client)
    job = (await client.post(
        "/api/workspaces/acme/exports/", json={"name": "Sheet", "format": "xlsx"})).json()
    assert job["status"] == "pending"
    await _run_job(db_session, job["id"])

    row = await db_session.get(ExportJob, job["id"])
    assert row.status == EExportStatus.READY
    assert row.content is not None
    wb = load_workbook(BytesIO(base64.b64decode(row.content)))
    ws = wb.active
    assert [c.value for c in next(ws.iter_rows(max_row=1))] == [
        "id", "title", "status", "priority",
    ]


async def test_xlsx_download_bytes_and_content_type(app, client, db_session):
    await setup(client)
    job = (await client.post(
        "/api/workspaces/acme/exports/", json={"name": "Sheet", "format": "xlsx"})).json()
    await _run_job(db_session, job["id"])

    dl = await client.get(f"/api/workspaces/acme/exports/{job['id']}/download/")
    assert dl.status_code == 200
    assert dl.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = load_workbook(BytesIO(dl.content))
    assert wb.active["A1"].value == "id"


async def test_invalid_format_400(app, client):
    await setup(client)
    res = await client.post("/api/workspaces/acme/exports/", json={"name": "X", "format": "xml"})
    assert res.status_code == 400


async def test_download_unknown_404(app, client):
    await setup(client)
    res = await client.get("/api/workspaces/acme/exports/nope/download/")
    assert res.status_code == 404


async def test_download_not_ready_404(app, client, db_session):
    """A still-PENDING job has no content yet — download 404s (existing rule),
    it doesn't error just because the job exists. (The background render actually
    runs inline within the httpx ASGITransport request cycle — see run_export_in_
    fresh_session's docstring — so a genuinely-pending row is forced here rather
    than relied on to still be pending by the time this test's GET runs.)"""
    await setup(client)
    job = (await client.post(
        "/api/workspaces/acme/exports/", json={"name": "Backlog", "format": "csv"})).json()
    row = await db_session.get(ExportJob, job["id"])
    row.status = EExportStatus.PENDING
    row.content = None
    await db_session.commit()

    res = await client.get(f"/api/workspaces/acme/exports/{job['id']}/download/")
    assert res.status_code == 404


async def test_expired_export_download_410(app, client, db_session):
    await setup(client)
    job = (await client.post(
        "/api/workspaces/acme/exports/", json={"name": "Backlog", "format": "csv"})).json()
    await _run_job(db_session, job["id"])

    row = await db_session.get(ExportJob, job["id"])
    row.expires_at = utcnow() - timedelta(days=1)
    await db_session.commit()

    res = await client.get(f"/api/workspaces/acme/exports/{job['id']}/download/")
    assert res.status_code == 410


async def test_workspace_lookup_failure_does_not_propagate_and_job_stays_ready(
    app, client, db_session, monkeypatch
):
    """The post-render Workspace lookup (for the export.ready notification) used to
    sit outside any try/except — a DB blip there raised out of the background task
    even though the job had already committed READY. Guard it so the task never
    re-raises and the job still ends terminal (READY), not stuck/reverted."""
    from sqlalchemy.ext.asyncio import AsyncSession as SAAsyncSession

    from gigsnano_api.models import Workspace

    original_get = SAAsyncSession.get

    async def _flaky_get(self, entity, ident, *a, **kw):
        if entity is Workspace:
            raise RuntimeError("boom")
        return await original_get(self, entity, ident, *a, **kw)

    monkeypatch.setattr(SAAsyncSession, "get", _flaky_get)

    await setup(client)
    job = (await client.post(
        "/api/workspaces/acme/exports/", json={"name": "Backlog", "format": "csv"})).json()

    row = await db_session.get(ExportJob, job["id"])
    assert row.status == EExportStatus.READY


async def test_failed_render_flips_job_to_failed(app, client, db_session, monkeypatch):
    """The background render runs inline within the request (see run_export_in_
    fresh_session's docstring), so patch _render before the POST, not after."""
    import gigsnano_api.features.exports.router as export_router

    def _boom(stories, fmt):
        raise RuntimeError("boom")

    monkeypatch.setattr(export_router, "_render", _boom)

    await setup(client)
    job = (await client.post(
        "/api/workspaces/acme/exports/", json={"name": "Backlog", "format": "csv"})).json()

    row = await db_session.get(ExportJob, job["id"])
    assert row.status == EExportStatus.FAILED
    assert row.content is None
